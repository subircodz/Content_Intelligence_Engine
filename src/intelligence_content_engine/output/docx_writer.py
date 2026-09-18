"""DOCX rendering for generated article drafts."""
from __future__ import annotations
import re
from pathlib import Path
from typing import Optional
from docx import Document

def safe_filename(title: str) -> str:
    value = re.sub(r"[^a-zA-Z0-9]+", "-", str(title).strip().lower())
    value = re.sub(r"^-+|-+$", "", value)
    return value or "untitled-article"

def _add_inline_markdown(paragraph, text: str) -> None:
    pattern = re.compile(r"(\\*\\*[^*]+\\*\\*|\\*[^*]+\\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2]); run.bold = True
        else:
            run = paragraph.add_run(token[1:-1]); run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])

def markdown_to_docx(article: str, title: str) -> Document:
    document = Document()
    document.core_properties.title = title
    document.add_paragraph(title, style="Title")
    lines = article.replace("\\r\\n", "\\n").replace("\\r", "\\n").split("\\n")
    for raw_line in lines:
        line = raw_line.strip()
        if not line: continue
        heading = re.match(r"^(#{1,6})\\s+(.+)$", line)
        if heading:
            p = document.add_paragraph(style=f"Heading {min(len(heading.group(1)), 6)}")
            _add_inline_markdown(p, heading.group(2).strip()); continue
        bullet = re.match(r"^[-*+]\\s+(.+)$", line)
        if bullet:
            p = document.add_paragraph(style="List Bullet")
            _add_inline_markdown(p, bullet.group(1)); continue
        numbered = re.match(r"^\\d+[.)]\\s+(.+)$", line)
        if numbered:
            p = document.add_paragraph(style="List Number")
            _add_inline_markdown(p, numbered.group(1)); continue
        p = document.add_paragraph(); _add_inline_markdown(p, line)
    return document

def save_article_docx(article: str, title: str, output_dir: str = "output", competitor_analysis: Optional[object] = None) -> Optional[Path]:
    if not article or not article.strip(): return None
    try:
        directory = Path(output_dir); directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{safe_filename(title)}.docx"
        document = markdown_to_docx(article, title)
        if competitor_analysis is not None:
            coverage = getattr(competitor_analysis, "coverage", None); gaps = getattr(competitor_analysis, "gaps", None)
            if coverage is not None and gaps is not None:
                document.add_heading("Editorial Planning Notes", level=1)
                p = document.add_paragraph(); p.add_run("Market coverage: ").bold = True; p.add_run(str(getattr(coverage, "status", "UNKNOWN")))
                p = document.add_paragraph(); p.add_run("Opportunity type: ").bold = True; p.add_run(str(getattr(coverage, "opportunity_type", "UNKNOWN")))
                for label, values in (("Missing topics", getattr(gaps, "missing_topics", [])), ("Missing questions", getattr(gaps, "missing_questions", [])), ("Missing entities", getattr(gaps, "missing_entities", [])), ("Missing comparisons", getattr(gaps, "missing_comparisons", [])), ("Recommended angles", getattr(gaps, "missing_angles", []))):
                    if values:
                        document.add_heading(label, level=2)
                        for value in values: document.add_paragraph(str(value), style="List Bullet")
        document.save(path); return path
    except Exception:
        return None
